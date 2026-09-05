"""
PRD Generator - Professional Edition
=====================================
Generates world-class PRDs like top PMs at FAANG companies.

Flow: Input -> Generate -> Refine -> Export

Features:
- FAANG-quality PRD output, section by section
- Strict Mode: never invents metrics that weren't given
- Per-section AI refinement ("regenerate with feedback")
- Status, tags, search/filter, version history, duplicate
- Multiple export formats (Markdown, HTML, Word) + shareable link
- Friendly, specific error messages instead of raw tracebacks

Run:
    pip install -r requirements.txt
    streamlit run spec_to_prd_app.py
"""

import json
import os
import re
import time
import base64
import hashlib
from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, List
from io import BytesIO
import streamlit as st

# =============================================================================
# IMPORTS WITH FALLBACKS
# =============================================================================

try:
    from groq import Groq
    GROQ_AVAILABLE = True
except ImportError:
    GROQ_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# =============================================================================
# CONFIGURATION
# =============================================================================

STORAGE_DIR = Path.home() / ".prd_generator"
PRDS_FILE = STORAGE_DIR / "prds.json"

MODEL_LABELS = {
    "openai/gpt-oss-120b": "GPT-OSS 120B — best quality (recommended)",
    "openai/gpt-oss-20b": "GPT-OSS 20B — fastest",
}
MODEL_IDS = list(MODEL_LABELS.keys())

STATUSES = ["Draft", "In Review", "Approved"]
STATUS_COLORS = {"Draft": "#6b7280", "In Review": "#b45309", "Approved": "#15803d"}

SECTION_ICONS = {
    "overview": "🧭", "problem": "🚨", "goals": "🎯", "users": "👥",
    "solution": "💡", "requirements": "✅", "scope": "🗓️", "risks": "⚠️",
}

# =============================================================================
# WORLD-CLASS PRD SECTIONS
# =============================================================================

SECTIONS = [
    {
        "id": "overview",
        "name": "Overview",
        "prompt": """Write a crisp Overview section (3-4 sentences max).

Include:
- One sentence: What we're building
- One sentence: Why now (urgency/opportunity)
- One sentence: Expected business impact

Be specific. No fluff. Every word should matter."""
    },
    {
        "id": "problem",
        "name": "Problem Statement",
        "prompt": """Write a compelling Problem Statement.

Structure:
1. **The Problem** (2-3 sentences): What's broken? Be specific.
2. **Who's Affected**: Which users/segments experience this?
3. **Current Workarounds**: How do users solve this today?
4. **Cost of Inaction**: What happens if we don't solve this?

Bad: "Users find it hard to export data"
Good: "Enterprise users (45% of revenue) cannot export to Excel, causing 3+ support tickets/week and blocking 2 active deals worth $50K ARR" """
    },
    {
        "id": "goals",
        "name": "Goals & Success Metrics",
        "prompt": """Write Goals & Success Metrics that an executive would approve.

Structure:
**Primary Goal**: One sentence, measurable

**Success Metrics** (table format):
| Metric | Current | Target | Timeline |
|--------|---------|--------|----------|
| Example: Support tickets for exports | 12/week | <3/week | 30 days post-launch |

**Non-Goals** (equally important):
- What we're explicitly NOT trying to achieve
- What's out of scope for this iteration

Avoid vanity metrics.
Bad: "Improve user satisfaction"
Good: "Reduce export-related support tickets from 12/week to <3/week within 30 days" """
    },
    {
        "id": "users",
        "name": "User Stories & Personas",
        "prompt": """Write User Stories that engineers can build from.

**Primary Persona**:
Name, role, context (1-2 sentences)

**User Stories** (use this exact format):
1. As a [specific user], I want to [specific action], so that [specific outcome].
   - Acceptance criteria: [How we know it's done]

2. As a [specific user], I want to [specific action], so that [specific outcome].
   - Acceptance criteria: [How we know it's done]

Include 3-5 user stories covering the main use cases.

Bad: "As a user, I want to export data"
Good: "As a finance analyst, I want to export quarterly data to Excel with formulas preserved, so that I can run my existing pivot table reports without manual reformatting. Acceptance: Excel file opens with formulas working, <30 sec for 100K rows" """
    },
    {
        "id": "solution",
        "name": "Proposed Solution",
        "prompt": """Write the Proposed Solution section.

Structure:
**Solution Summary**: 2-3 sentences describing the approach

**Key Capabilities**:
1. [Capability 1]: Brief description
2. [Capability 2]: Brief description
3. [Capability 3]: Brief description

**Why This Approach**:
- Alternative considered: [X] — Why rejected: [reason]
- Alternative considered: [Y] — Why rejected: [reason]
- Chosen approach: [Z] — Why: [compelling reason]

**Technical Approach** (high-level):
Brief description of how this will be built (1-2 sentences)

Be specific about what we're building, not vague aspirations."""
    },
    {
        "id": "requirements",
        "name": "Functional Requirements",
        "prompt": """Write Functional Requirements that an engineer can build from tomorrow.

**P0 - Must Have** (launch blockers):
| ID | Requirement | Details | Acceptance Criteria |
|----|-------------|---------|---------------------|
| R1 | [Requirement] | [Specifics] | [How to verify] |

**P1 - Should Have** (fast follow):
| ID | Requirement | Details | Acceptance Criteria |
|----|-------------|---------|---------------------|

**P2 - Nice to Have** (future):
| ID | Requirement | Details | Acceptance Criteria |
|----|-------------|---------|---------------------|

Be extremely specific. Engineers should not need to ask clarifying questions.
Bad: "Support large files"
Good: "Support files up to 1M rows, export completes in <60 seconds, progress indicator shown" """
    },
    {
        "id": "scope",
        "name": "Scope & Timeline",
        "prompt": """Write Scope & Timeline section.

**In Scope** (what we WILL build):
- [Specific deliverable 1]
- [Specific deliverable 2]
- [Specific deliverable 3]

**Out of Scope** (what we will NOT build):
- [Explicit exclusion 1] — Why: [reason]
- [Explicit exclusion 2] — Why: [reason]

**Dependencies**:
- [Team/System]: [What we need from them]

**Timeline**:
| Milestone | Date | Owner |
|-----------|------|-------|
| Design complete | Week 1 | Design |
| Backend API ready | Week 2 | Backend |
| Frontend complete | Week 3 | Frontend |
| QA & Bug fixes | Week 4 | QA |
| Launch | Week 5 | PM |

Be explicit about what's NOT included. This prevents scope creep."""
    },
    {
        "id": "risks",
        "name": "Risks & Mitigations",
        "prompt": """Write Risks & Mitigations like a senior PM who's shipped before.

| Risk | Likelihood | Impact | Mitigation |
|------|------------|--------|------------|
| [Specific risk] | High/Med/Low | High/Med/Low | [Specific action to reduce risk] |

Include:
- Technical risks (performance, scale, integration)
- Business risks (adoption, competition, timing)
- Resource risks (dependencies, availability)

Be honest about what could go wrong. Leadership respects realism.
Bad: "There might be technical challenges"
Good: "Export >500K rows may timeout on current infrastructure. Mitigation: Implement async processing with email notification, already validated approach with Platform team" """
    },
]

# =============================================================================
# PROMPT TEMPLATES
# =============================================================================

DATA_INTEGRITY_STRICT = """DATA INTEGRITY RULE (do not break this):
If the stakeholder input below does not contain a specific number, date, or metric you need,
do NOT invent one. Write `[NEEDS INPUT: <what's missing>]` in its place and move on.
Never fabricate statistics, dollar amounts, percentages, or dates that were not given or
clearly implied by the input."""

DATA_INTEGRITY_LOOSE = """DATA INTEGRITY RULE:
If the stakeholder input below does not contain a specific number, use a clearly-labeled
reasonable estimate, e.g. "(est.)", so the reader knows it wasn't provided."""

SECTION_PROMPT_TEMPLATE = """You are a Senior Product Manager at a top tech company (Google/Stripe/Airbnb level).

You're writing a PRD that will be reviewed by:
- Engineering leads (who need to estimate and build)
- Design leads (who need to understand the experience)
- Executive sponsors (who need to approve resources)

Your PRDs are known for:
- Crystal clarity (no ambiguity)
- Specific metrics (real numbers)
- Actionable requirements (engineers can start tomorrow)
- Honest risks (no surprises later)

{data_integrity_rule}

FEATURE CONTEXT:
===============
Feature Name: {feature_name}

Raw Input from Stakeholder:
{raw_input}

YOUR TASK:
==========
Write the "{section_name}" section of this PRD.

{section_specific_prompt}

QUALITY BAR:
============
- Every sentence must add value (no filler)
- Be specific (numbers, names, dates) but only when the input supports it
- Be honest (include challenges, not just benefits)
- Be actionable (reader knows exactly what to do next)

Write ONLY the section content. No preamble, no "Here's the section", just the content."""

REGENERATE_TEMPLATE = """You previously wrote the "{section_name}" section of a PRD for "{feature_name}":

---
{previous_content}
---

The author wants this specific change applied: "{feedback}"

Rewrite the section incorporating that feedback. Keep the same structure/format
conventions (headings, tables, lists) as the original. Output ONLY the revised
section content, no preamble."""

# =============================================================================
# STORAGE
# =============================================================================

def init_storage():
    STORAGE_DIR.mkdir(exist_ok=True)
    if not PRDS_FILE.exists():
        PRDS_FILE.write_text("[]")

def load_prds() -> List[Dict]:
    init_storage()
    try:
        data = PRDS_FILE.read_text()
        return json.loads(data) if data.strip() else []
    except Exception:
        return []

def save_prds_to_disk(prds: List[Dict]):
    init_storage()
    PRDS_FILE.write_text(json.dumps(prds, indent=2, default=str))

def save_prd(prd: Dict) -> str:
    prds = load_prds()

    is_new = "id" not in prd or not prd["id"]
    if is_new:
        prd["id"] = hashlib.md5(f"{time.time()}".encode()).hexdigest()[:8]
        prd["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        prd.setdefault("status", "Draft")
        prd.setdefault("tags", [])
        prd.setdefault("history", [])

    existing_idx = next((i for i, p in enumerate(prds) if p.get("id") == prd["id"]), None)
    if existing_idx is not None:
        old = prds[existing_idx]
        if old.get("content") and old.get("content") != prd.get("content"):
            history = old.get("history", [])
            history.insert(0, {"content": old["content"], "saved_at": old.get("updated_at", "")})
            prd["history"] = history[:10]
        prds[existing_idx] = prd
    else:
        prds.insert(0, prd)

    prd["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    save_prds_to_disk(prds)
    return prd["id"]

def delete_prd(prd_id: str):
    prds = load_prds()
    prds = [p for p in prds if p.get("id") != prd_id]
    save_prds_to_disk(prds)

def duplicate_prd(prd: Dict) -> Dict:
    new_prd = json.loads(json.dumps(prd))
    new_prd["id"] = None
    new_prd["name"] = f"{prd.get('name', 'Untitled')} (copy)"
    new_prd["status"] = "Draft"
    new_prd["history"] = []
    save_prd(new_prd)
    return new_prd

# =============================================================================
# SHARING
# =============================================================================

def encode_for_sharing(prd: Dict) -> str:
    share_data = {"n": prd.get("name", ""), "c": prd.get("content", "")}
    json_str = json.dumps(share_data, separators=(',', ':'))
    return base64.urlsafe_b64encode(json_str.encode()).decode()

def decode_shared(encoded: str) -> Optional[Dict]:
    try:
        json_str = base64.urlsafe_b64decode(encoded.encode()).decode()
        data = json.loads(json_str)
        return {"name": data.get("n", "Shared PRD"), "content": data.get("c", "")}
    except Exception:
        return None

def get_base_url() -> str:
    try:
        headers = st.context.headers
        host = headers.get("Host")
        if host:
            scheme = "http" if host.startswith("localhost") or host.startswith("127.0.0.1") else "https"
            return f"{scheme}://{host}"
    except Exception:
        pass
    return "http://localhost:8501"

def render_share_button(prd: Dict, key_prefix: str):
    if st.button("🔗 Share", key=f"{key_prefix}_share", use_container_width=True):
        encoded = encode_for_sharing(prd)
        link = f"{get_base_url()}/?shared={encoded}"
        if len(link) > 6000:
            st.warning(
                "This PRD is large, so the share link is long. Some chat apps/SMS "
                "truncate long links — exporting and sending the file is more reliable."
            )
        st.code(link, language=None)

# =============================================================================
# MARKDOWN HELPERS (shared by HTML + Word export)
# =============================================================================

def _is_table_separator(stripped: str) -> bool:
    return stripped.replace(" ", "").replace("|", "").replace("-", "") == ""

def _parse_table_row(stripped: str) -> List[str]:
    return [c.strip() for c in stripped.split("|")[1:-1]]

# =============================================================================
# EXPORT FUNCTIONS
# =============================================================================

def export_markdown(prd: Dict) -> str:
    return prd.get("content", "")

def export_html(prd: Dict) -> str:
    name = prd.get("name", "PRD")
    content = prd.get("content", "")

    lines = []
    in_table = False
    table_first_row = False

    for line in content.split("\n"):
        stripped = line.strip()

        if stripped.startswith("|"):
            if not in_table:
                lines.append("<table style='border-collapse: collapse; width: 100%; margin: 15px 0; font-size: 14px;'>")
                in_table = True
                table_first_row = True
            if _is_table_separator(stripped):
                continue
            cells = _parse_table_row(stripped)
            if table_first_row:
                lines.append("<tr>" + "".join(f"<th style='border: 1px solid #ddd; padding: 10px; background: #f5f5f5; text-align: left;'>{c}</th>" for c in cells) + "</tr>")
                table_first_row = False
            else:
                lines.append("<tr>" + "".join(f"<td style='border: 1px solid #ddd; padding: 10px;'>{c}</td>" for c in cells) + "</tr>")
            continue
        elif in_table:
            lines.append("</table>")
            in_table = False

        if stripped.startswith("### "):
            lines.append(f"<h3 style='color: #333; margin-top: 25px; font-size: 16px;'>{stripped[4:]}</h3>")
        elif stripped.startswith("## "):
            lines.append(f"<h2 style='color: #1a56db; margin-top: 35px; padding-bottom: 10px; border-bottom: 2px solid #1a56db; font-size: 20px;'>{stripped[3:]}</h2>")
        elif stripped.startswith("# "):
            lines.append(f"<h1 style='color: #111; font-size: 28px;'>{stripped[2:]}</h1>")
        elif stripped.startswith("- "):
            lines.append(f"<li style='margin: 8px 0 8px 20px;'>{stripped[2:]}</li>")
        elif stripped == "---":
            lines.append("<hr style='margin: 25px 0; border: none; border-top: 1px solid #e5e7eb;'>")
        elif stripped:
            formatted = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', stripped)
            lines.append(f"<p style='margin: 12px 0; line-height: 1.7;'>{formatted}</p>")

    if in_table:
        lines.append("</table>")

    html_content = "\n".join(lines)

    return f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{name}</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', sans-serif;
            max-width: 850px;
            margin: 50px auto;
            padding: 30px;
            line-height: 1.6;
            color: #333;
            background: #fff;
        }}
        .header {{
            border-bottom: 1px solid #e5e7eb;
            padding-bottom: 20px;
            margin-bottom: 30px;
        }}
        .meta {{
            color: #6b7280;
            font-size: 13px;
        }}
    </style>
</head>
<body>
    <div class="header">
        <div class="meta">PRD Document &bull; Generated {datetime.now().strftime("%Y-%m-%d")}</div>
    </div>
    {html_content}
</body>
</html>"""

def export_docx(prd: Dict) -> Optional[bytes]:
    if not DOCX_AVAILABLE:
        return None

    doc = Document()
    doc.add_heading(prd.get("name", "PRD"), 0)

    meta = doc.add_paragraph()
    meta.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d')}").italic = True
    doc.add_paragraph()

    content = prd.get("content", "")
    in_table = False
    table_data = []

    def flush_table():
        if table_data and table_data[0]:
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            table.style = 'Table Grid'
            for i, row in enumerate(table_data):
                for j, cell in enumerate(row):
                    if j < len(table.rows[i].cells):
                        table.rows[i].cells[j].text = cell
                        if i == 0:
                            for para in table.rows[i].cells[j].paragraphs:
                                for run in para.runs:
                                    run.bold = True
            doc.add_paragraph()
        table_data.clear()

    for line in content.split("\n"):
        stripped = line.strip()

        if not stripped:
            continue

        if stripped.startswith("|"):
            in_table = True
            if _is_table_separator(stripped):
                continue
            cells = _parse_table_row(stripped)
            if cells:
                table_data.append(cells)
            continue
        elif in_table:
            flush_table()
            in_table = False

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif stripped == "---":
            pass
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.+?\*\*)', stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table()

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# =============================================================================
# AI CLIENT
# =============================================================================

class AIClient:
    def __init__(self, api_key: str, model: str):
        if not GROQ_AVAILABLE:
            raise RuntimeError("Install groq: pip install groq")
        self.client = Groq(api_key=api_key)
        self.model = model

    def generate(self, prompt: str, max_retries: int = 3) -> str:
        last_err = None
        for attempt in range(max_retries):
            try:
                response = self.client.chat.completions.create(
                    model=self.model,
                    messages=[{"role": "user", "content": prompt}],
                    max_tokens=4096,
                    temperature=0.3
                )
                return response.choices[0].message.content
            except Exception as e:
                err_name = type(e).__name__
                status = getattr(e, "status_code", None)
                if err_name == "AuthenticationError" or status == 401:
                    raise RuntimeError("Invalid Groq API key. Double-check it in the sidebar under Settings.") from e
                last_err = e
                if attempt < max_retries - 1:
                    time.sleep(min(2 ** attempt, 8))
                    continue
                if err_name == "RateLimitError" or status == 429:
                    raise RuntimeError("Groq rate limit reached. Wait a minute, or switch to a different model in Settings.") from e
        raise RuntimeError(f"AI provider error: {last_err}")

# =============================================================================
# STREAMLIT APP
# =============================================================================

st.set_page_config(
    page_title="PRD Generator",
    page_icon="📝",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
    .block-container { padding-top: 2rem; max-width: 1100px; }
    div.stButton > button { border-radius: 8px; }
    .prd-card {
        border: 1px solid rgba(120,120,120,0.2);
        border-radius: 10px;
        padding: 14px 18px;
        margin-bottom: 10px;
    }
    .badge {
        display: inline-block;
        padding: 2px 10px;
        border-radius: 999px;
        font-size: 12px;
        font-weight: 600;
        color: white;
    }
    .hero-tile {
        border: 1px solid rgba(120,120,120,0.2);
        border-radius: 10px;
        padding: 16px 20px;
        text-align: center;
    }
    .hero-number { font-size: 30px; font-weight: 700; }
    .hero-label { font-size: 13px; opacity: 0.7; }
</style>
""", unsafe_allow_html=True)

def status_badge_html(status: str) -> str:
    color = STATUS_COLORS.get(status, "#6b7280")
    return f"<span class='badge' style='background:{color}'>{status}</span>"

# Session state initialization
if "page" not in st.session_state:
    st.session_state.page = "home"
if "current_prd" not in st.session_state:
    st.session_state.current_prd = {}
if "sections_content" not in st.session_state:
    st.session_state.sections_content = {}
if "generation_complete" not in st.session_state:
    st.session_state.generation_complete = False
if "editing_section" not in st.session_state:
    st.session_state.editing_section = None
if "api_key" not in st.session_state:
    st.session_state.api_key = ""
if "model" not in st.session_state:
    st.session_state.model = MODEL_IDS[0]
if "strict_mode" not in st.session_state:
    st.session_state.strict_mode = True
if "pending_delete" not in st.session_state:
    st.session_state.pending_delete = None

# Check for shared PRD in URL
query_params = st.query_params
if "shared" in query_params and "shared_prd" not in st.session_state:
    shared = decode_shared(query_params["shared"])
    if shared:
        st.session_state.shared_prd = shared
        st.session_state.page = "view_shared"

# =============================================================================
# SIDEBAR
# =============================================================================

with st.sidebar:
    st.title("📝 PRD Generator")
    st.caption("FAANG-quality PRDs in minutes, not days")

    st.divider()

    if st.button("🏠 Home", use_container_width=True):
        st.session_state.page = "home"
        st.rerun()

    if st.button("✨ New PRD", use_container_width=True):
        st.session_state.page = "new"
        st.session_state.current_prd = {}
        st.session_state.sections_content = {}
        st.session_state.generation_complete = False
        st.session_state.editing_section = None
        st.rerun()

    if st.button("📚 My PRDs", use_container_width=True):
        st.session_state.page = "history"
        st.rerun()

    st.divider()
    st.subheader("⚙️ Settings")

    api_key_loaded = False
    if hasattr(st, 'secrets') and "GROQ_API_KEY" in st.secrets:
        st.session_state.api_key = st.secrets["GROQ_API_KEY"]
        api_key_loaded = True
    elif os.getenv("GROQ_API_KEY"):
        st.session_state.api_key = os.getenv("GROQ_API_KEY")
        api_key_loaded = True

    if not api_key_loaded:
        entered_key = st.text_input(
            "Groq API key",
            value=st.session_state.api_key,
            type="password",
            help="Get a free key at console.groq.com/keys. Only kept in this browser session, never saved to disk.",
        )
        if entered_key:
            st.session_state.api_key = entered_key
            api_key_loaded = True

    model = st.selectbox(
        "Model",
        MODEL_IDS,
        index=MODEL_IDS.index(st.session_state.model) if st.session_state.model in MODEL_IDS else 0,
        format_func=lambda m: MODEL_LABELS.get(m, m),
    )
    st.session_state.model = model

    st.session_state.strict_mode = st.toggle(
        "Strict mode (no invented numbers)",
        value=st.session_state.strict_mode,
        help="When on, the AI writes [NEEDS INPUT: ...] instead of making up metrics you didn't provide.",
    )

    if api_key_loaded:
        st.success("✅ API connected")
    else:
        st.error("❌ No API key")
        st.caption("Paste a key above, or add GROQ_API_KEY to Streamlit secrets for a hidden shared key.")

    st.divider()
    st.caption("Built by Prem Dutta")

# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def get_ai_client():
    if not st.session_state.api_key:
        st.error("API key not configured. Paste a Groq API key in the sidebar under Settings.")
        st.stop()
    return AIClient(st.session_state.api_key, st.session_state.model)

def build_prd_content(name: str, sections: Dict) -> str:
    lines = [f"# PRD: {name}", "", "---", ""]
    for section in SECTIONS:
        content = sections.get(section["id"], "")
        if content:
            lines.extend([f"## {section['name']}", "", content, ""])
    return "\n".join(lines)

def parse_sections_from_content(content: str) -> Dict[str, str]:
    parsed = {}
    for sec in SECTIONS:
        marker = f"## {sec['name']}"
        if marker not in content:
            continue
        start = content.index(marker) + len(marker)
        end = len(content)
        for next_sec in SECTIONS:
            next_marker = f"## {next_sec['name']}"
            if next_marker in content[start:]:
                end = start + content[start:].index(next_marker)
                break
        parsed[sec["id"]] = content[start:end].strip()
    return parsed

def word_count(text: str) -> int:
    return len(text.split())

def data_integrity_rule() -> str:
    return DATA_INTEGRITY_STRICT if st.session_state.strict_mode else DATA_INTEGRITY_LOOSE

# =============================================================================
# PAGE: HOME
# =============================================================================

if st.session_state.page == "home":
    st.title("📝 PRD Generator")
    st.markdown("Generate professional PRDs like a Senior PM at Google, Stripe, or Airbnb.")

    st.divider()

    prds = load_prds()
    week_cutoff = datetime.now().timestamp() - 7 * 86400
    def _ts(p):
        try:
            return datetime.strptime(p.get("created_at", ""), "%Y-%m-%d %H:%M").timestamp()
        except Exception:
            return 0
    this_week = sum(1 for p in prds if _ts(p) >= week_cutoff)
    approved = sum(1 for p in prds if p.get("status") == "Approved")

    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown(f"<div class='hero-tile'><div class='hero-number'>{len(prds)}</div><div class='hero-label'>Total PRDs</div></div>", unsafe_allow_html=True)
    with c2:
        st.markdown(f"<div class='hero-tile'><div class='hero-number'>{this_week}</div><div class='hero-label'>Created this week</div></div>", unsafe_allow_html=True)
    with c3:
        st.markdown(f"<div class='hero-tile'><div class='hero-number'>{approved}</div><div class='hero-label'>Approved</div></div>", unsafe_allow_html=True)

    st.divider()

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("🚀 Get Started")
        st.markdown("""
        **Simple process:**
        1. Paste your rough notes or requirements
        2. Get a professional PRD instantly
        3. Refine sections with AI feedback
        4. Export or share
        """)

        if st.button("✨ Create New PRD", type="primary", use_container_width=True):
            st.session_state.page = "new"
            st.session_state.current_prd = {}
            st.session_state.sections_content = {}
            st.session_state.generation_complete = False
            st.rerun()

    with col2:
        st.subheader("📚 Recent PRDs")

        if prds:
            for prd in prds[:5]:
                with st.container():
                    st.markdown(
                        f"<div class='prd-card'><b>{prd.get('name', 'Untitled')}</b> {status_badge_html(prd.get('status', 'Draft'))}"
                        f"<div style='font-size:12px;opacity:0.6;margin-top:4px;'>{prd.get('created_at', '')[:10]}</div></div>",
                        unsafe_allow_html=True,
                    )
                    if st.button("Open", key=f"home_open_{prd.get('id')}"):
                        st.session_state.current_prd = prd
                        st.session_state.page = "view_prd"
                        st.rerun()
        else:
            st.info("No PRDs yet. Create your first one!")

# =============================================================================
# PAGE: NEW PRD
# =============================================================================

elif st.session_state.page == "new":

    if not st.session_state.generation_complete:
        st.title("✨ Create New PRD")
        st.markdown("Paste your rough notes below. The more context you provide, the better the PRD.")

        prd_name = st.text_input(
            "PRD Name",
            value=st.session_state.current_prd.get("name", ""),
            placeholder="e.g., Multi-Format Data Export"
        )

        col_a, col_b = st.columns(2)
        with col_a:
            prd_status = st.selectbox("Status", STATUSES, index=0)
        with col_b:
            prd_tags = st.text_input("Tags (comma separated)", placeholder="growth, mobile, Q1")

        raw_input = st.text_area(
            "Your Notes / Requirements",
            value=st.session_state.current_prd.get("raw_input", ""),
            height=300,
            placeholder="""Paste everything you have about this feature...

Example:
Feature: Data export in multiple formats

Problem:
- Users can only export to CSV
- Enterprise customers need Excel and PDF
- Lost 2 deals last quarter ($50K ARR)
- 10+ support tickets/week about exports

Users:
- Enterprise finance teams (need Excel)
- Compliance teams (need PDF reports)
- Data analysts (need large dataset support)

Solution ideas:
- Excel export with formatting
- PDF export for reports
- Support up to 1M rows

Timeline: Q1"""
        )

        with st.expander("💡 Tips for better output"):
            st.markdown("""
            Include as much context as you can:
            - **Problem**: What's broken? Include numbers if available.
            - **Users**: Who's affected? Which segments?
            - **Solution**: What are you thinking of building?
            - **Goals**: How will you measure success?
            - **Constraints**: Timeline, technical, dependencies

            With **Strict Mode** on (sidebar), the AI will never invent metrics you
            didn't give it — it'll flag them as `[NEEDS INPUT: ...]` so you can fill
            them in yourself.
            """)

        col1, col2 = st.columns([3, 1])
        with col2:
            if st.button("🚀 Generate PRD", type="primary", use_container_width=True):
                if not prd_name.strip():
                    st.error("Please enter a PRD name")
                elif not raw_input.strip():
                    st.error("Please enter your notes/requirements")
                elif not st.session_state.api_key:
                    st.error("API key not configured. Paste a Groq key in the sidebar.")
                else:
                    st.session_state.current_prd["name"] = prd_name.strip()
                    st.session_state.current_prd["raw_input"] = raw_input.strip()
                    st.session_state.current_prd["status"] = prd_status
                    st.session_state.current_prd["tags"] = [t.strip() for t in prd_tags.split(",") if t.strip()]

                    try:
                        client = get_ai_client()
                        sections = {}

                        with st.status("Generating PRD...", expanded=True) as status_box:
                            for section in SECTIONS:
                                status_box.write(f"{SECTION_ICONS.get(section['id'], '📄')} Writing: {section['name']}...")
                                prompt = SECTION_PROMPT_TEMPLATE.format(
                                    data_integrity_rule=data_integrity_rule(),
                                    feature_name=prd_name.strip(),
                                    raw_input=raw_input.strip(),
                                    section_name=section["name"],
                                    section_specific_prompt=section["prompt"]
                                )
                                content = client.generate(prompt)
                                sections[section["id"]] = content
                                status_box.write(f"✅ {section['name']} done")

                            status_box.update(label="PRD generated", state="complete", expanded=False)

                        st.session_state.sections_content = sections
                        st.session_state.current_prd["content"] = build_prd_content(prd_name, sections)
                        st.session_state.generation_complete = True
                        save_prd(st.session_state.current_prd)
                        st.rerun()

                    except RuntimeError as e:
                        st.error(str(e))
                    except Exception as e:
                        st.error(f"Something went wrong: {e}")

    else:
        prd = st.session_state.current_prd

        col_t, col_s = st.columns([4, 2])
        with col_t:
            st.title(f"📄 {prd.get('name', 'PRD')}")
        with col_s:
            new_status = st.selectbox(
                "Status", STATUSES,
                index=STATUSES.index(prd.get("status", "Draft")) if prd.get("status") in STATUSES else 0,
                key="status_select_new",
                label_visibility="collapsed",
            )
            if new_status != prd.get("status"):
                prd["status"] = new_status
                save_prd(prd)
                st.session_state.current_prd = prd
                st.rerun()

        col1, col2, col3, col4 = st.columns(4)
        with col1:
            md = export_markdown(prd)
            st.download_button("📄 Markdown", md, f"{prd.get('name', 'prd')}.md", use_container_width=True)
        with col2:
            html = export_html(prd)
            st.download_button("🌐 HTML", html, f"{prd.get('name', 'prd')}.html", use_container_width=True)
        with col3:
            if DOCX_AVAILABLE:
                docx = export_docx(prd)
                if docx:
                    st.download_button("📝 Word", docx, f"{prd.get('name', 'prd')}.docx", use_container_width=True)
            else:
                st.button("📝 Word", disabled=True, help="pip install python-docx", use_container_width=True)
        with col4:
            render_share_button(prd, key_prefix="new")

        st.divider()

        for section in SECTIONS:
            section_id = section["id"]
            section_name = section["name"]
            content = st.session_state.sections_content.get(section_id, "")

            with st.container(border=True):
                col1, col2 = st.columns([6, 1])
                with col1:
                    icon = SECTION_ICONS.get(section_id, "📄")
                    st.markdown(f"#### {icon} {section_name}")
                    st.caption(f"{word_count(content)} words")
                with col2:
                    if st.session_state.editing_section != section_id:
                        if st.button("✏️", key=f"edit_{section_id}", help="Edit"):
                            st.session_state.editing_section = section_id
                            st.rerun()

                if st.session_state.editing_section == section_id:
                    edited = st.text_area(
                        f"Edit {section_name}",
                        value=content,
                        height=300,
                        key=f"textarea_{section_id}",
                        label_visibility="collapsed"
                    )

                    c1, c2, c3 = st.columns([1, 1, 4])
                    with c1:
                        if st.button("💾 Save", key=f"save_{section_id}"):
                            st.session_state.sections_content[section_id] = edited
                            st.session_state.current_prd["content"] = build_prd_content(
                                prd.get("name", "PRD"), st.session_state.sections_content
                            )
                            save_prd(st.session_state.current_prd)
                            st.session_state.editing_section = None
                            st.rerun()
                    with c2:
                        if st.button("❌ Cancel", key=f"cancel_{section_id}"):
                            st.session_state.editing_section = None
                            st.rerun()
                else:
                    st.markdown(content)

                with st.expander("🔄 Refine with AI"):
                    feedback = st.text_input(
                        "What should change?",
                        key=f"feedback_{section_id}",
                        placeholder="e.g., make the timeline more aggressive",
                    )
                    if st.button("Regenerate this section", key=f"regen_{section_id}", disabled=not feedback.strip()):
                        try:
                            client = get_ai_client()
                            with st.spinner(f"Refining {section_name}..."):
                                prompt = REGENERATE_TEMPLATE.format(
                                    section_name=section_name,
                                    feature_name=prd.get("name", "PRD"),
                                    previous_content=content,
                                    feedback=feedback.strip(),
                                )
                                new_content = client.generate(prompt)
                            st.session_state.sections_content[section_id] = new_content
                            st.session_state.current_prd["content"] = build_prd_content(
                                prd.get("name", "PRD"), st.session_state.sections_content
                            )
                            save_prd(st.session_state.current_prd)
                            st.rerun()
                        except RuntimeError as e:
                            st.error(str(e))
                        except Exception as e:
                            st.error(f"Something went wrong: {e}")

        st.divider()
        col1, col2 = st.columns(2)
        with col1:
            if st.button("📚 All PRDs", use_container_width=True):
                st.session_state.page = "history"
                st.rerun()
        with col2:
            if st.button("✨ New PRD", type="primary", use_container_width=True):
                st.session_state.current_prd = {}
                st.session_state.sections_content = {}
                st.session_state.generation_complete = False
                st.session_state.editing_section = None
                st.rerun()

# =============================================================================
# PAGE: HISTORY
# =============================================================================

elif st.session_state.page == "history":
    st.title("📚 My PRDs")

    prds = load_prds()

    if not prds:
        st.info("No PRDs yet.")
        if st.button("✨ Create First PRD", type="primary"):
            st.session_state.page = "new"
            st.rerun()
    else:
        col_search, col_status, col_sort = st.columns([3, 2, 2])
        with col_search:
            search = st.text_input("🔍 Search", placeholder="Search by name or tag...")
        with col_status:
            status_filter = st.multiselect("Status", STATUSES)
        with col_sort:
            sort_by = st.selectbox("Sort by", ["Last updated", "Created", "Name"])

        if search:
            s = search.lower()
            prds = [p for p in prds if s in p.get("name", "").lower() or any(s in t.lower() for t in p.get("tags", []))]
        if status_filter:
            prds = [p for p in prds if p.get("status", "Draft") in status_filter]

        if sort_by == "Last updated":
            prds.sort(key=lambda p: p.get("updated_at", ""), reverse=True)
        elif sort_by == "Created":
            prds.sort(key=lambda p: p.get("created_at", ""), reverse=True)
        else:
            prds.sort(key=lambda p: p.get("name", "").lower())

        st.divider()

        for prd in prds:
            pid = prd.get("id")
            with st.container(border=True):
                col1, col2, col3, col4, col5 = st.columns([4, 1, 1, 1, 1])

                with col1:
                    tags = ", ".join(prd.get("tags", []))
                    st.markdown(f"**{prd.get('name', 'Untitled')}** {status_badge_html(prd.get('status', 'Draft'))}", unsafe_allow_html=True)
                    caption = f"Updated: {prd.get('updated_at', prd.get('created_at', 'N/A'))[:16]}"
                    if tags:
                        caption += f" • {tags}"
                    st.caption(caption)

                with col2:
                    if st.button("Open", key=f"hist_open_{pid}"):
                        st.session_state.current_prd = prd
                        st.session_state.sections_content = parse_sections_from_content(prd.get("content", ""))
                        st.session_state.generation_complete = True
                        st.session_state.page = "new"
                        st.rerun()

                with col3:
                    if st.button("⧉", key=f"dup_{pid}", help="Duplicate"):
                        duplicate_prd(prd)
                        st.rerun()

                with col4:
                    md = export_markdown(prd)
                    st.download_button("📥", md, f"{prd.get('name', 'prd')}.md", key=f"dl_{pid}")

                with col5:
                    if st.session_state.pending_delete == pid:
                        if st.button("✅", key=f"confirm_del_{pid}", help="Confirm delete"):
                            delete_prd(pid)
                            st.session_state.pending_delete = None
                            st.rerun()
                    else:
                        if st.button("🗑️", key=f"del_{pid}", help="Delete"):
                            st.session_state.pending_delete = pid
                            st.rerun()

                if st.session_state.pending_delete == pid:
                    st.warning(f"Delete \"{prd.get('name', 'Untitled')}\" permanently? Click ✅ to confirm.")

# =============================================================================
# PAGE: VIEW SHARED
# =============================================================================

elif st.session_state.page == "view_shared":
    prd = st.session_state.get("shared_prd", {})

    st.title("📄 Shared PRD")
    st.info("This PRD was shared with you as a read-only link.")

    st.subheader(prd.get("name", "Shared PRD"))
    st.divider()
    st.markdown(prd.get("content", "No content"))
    st.divider()

    col1, col2 = st.columns(2)
    with col1:
        if st.button("💾 Save to My PRDs", type="primary", use_container_width=True):
            imported = {"name": f"{prd.get('name', 'Shared')} (imported)", "content": prd.get("content", "")}
            save_prd(imported)
            st.success("✅ Saved!")
            time.sleep(1)
            st.session_state.page = "history"
            st.rerun()
    with col2:
        md = export_markdown(prd)
        st.download_button("📥 Download", md, "shared_prd.md", use_container_width=True)

# =============================================================================
# PAGE: VIEW PRD
# =============================================================================

elif st.session_state.page == "view_prd":
    prd = st.session_state.current_prd

    if not prd:
        st.session_state.page = "history"
        st.rerun()

    col1, col2 = st.columns([5, 1])
    with col1:
        st.title(prd.get("name", "PRD"))
        st.markdown(status_badge_html(prd.get("status", "Draft")), unsafe_allow_html=True)
    with col2:
        if st.button("← Back"):
            st.session_state.page = "history"
            st.rerun()

    st.caption(f"Created: {prd.get('created_at', 'N/A')}")

    col1, col2, col3 = st.columns(3)
    with col1:
        md = export_markdown(prd)
        st.download_button("📄 Markdown", md, f"{prd.get('name', 'prd')}.md", use_container_width=True)
    with col2:
        html = export_html(prd)
        st.download_button("🌐 HTML", html, f"{prd.get('name', 'prd')}.html", use_container_width=True)
    with col3:
        render_share_button(prd, key_prefix="view")

    st.divider()
    st.markdown(prd.get("content", "No content"))

# =============================================================================
# FOOTER
# =============================================================================

st.markdown("---")
st.caption("PRD Generator • Built by Prem Dutta")
